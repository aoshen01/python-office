"""Tests for office.word module."""

import os
import tempfile

import pytest
from office.word import Word


@pytest.fixture
def tmp_dir():
    with tempfile.TemporaryDirectory() as d:
        yield d


def test_create_minimal(tmp_dir):
    path = os.path.join(tmp_dir, "test.docx")
    result = Word.create(path)
    assert os.path.isfile(result)


def test_create_with_title_and_paragraphs(tmp_dir):
    path = os.path.join(tmp_dir, "doc.docx")
    result = Word.create(
        path,
        title="Hello World",
        paragraphs=["First paragraph.", "Second paragraph."],
    )
    assert os.path.isfile(result)


def test_extract_text(tmp_dir):
    path = os.path.join(tmp_dir, "doc.docx")
    Word.create(path, paragraphs=["Line one", "Line two"])
    text = Word.extract_text(path)
    assert "Line one" in text
    assert "Line two" in text


def test_get_paragraphs(tmp_dir):
    path = os.path.join(tmp_dir, "doc.docx")
    Word.create(path, paragraphs=["Alpha", "Beta", "Gamma"])
    paras = Word.get_paragraphs(path)
    assert "Alpha" in paras
    assert "Beta" in paras
    assert "Gamma" in paras


def test_to_txt(tmp_dir):
    docx_path = os.path.join(tmp_dir, "doc.docx")
    Word.create(docx_path, paragraphs=["Export me"])
    txt_path = Word.to_txt(docx_path)
    assert txt_path.endswith(".txt")
    assert os.path.isfile(txt_path)
    with open(txt_path, encoding="utf-8") as f:
        content = f.read()
    assert "Export me" in content


def test_add_table(tmp_dir):
    docx_path = os.path.join(tmp_dir, "doc.docx")
    Word.create(docx_path, title="Table Doc")
    data = [["Name", "Score"], ["Alice", "95"], ["Bob", "87"]]
    result = Word.add_table(docx_path, data)
    assert os.path.isfile(result)
    # Verify the table is readable
    from docx import Document
    doc = Document(result)
    assert len(doc.tables) == 1
    assert doc.tables[0].cell(0, 0).text == "Name"
