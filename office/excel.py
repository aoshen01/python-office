"""Excel spreadsheet processing utilities."""

from __future__ import annotations

import csv
import os
from pathlib import Path
from typing import Any, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


class Excel:
    """Utilities for reading, writing, and converting Excel (.xlsx) files."""

    # ------------------------------------------------------------------
    # Reading
    # ------------------------------------------------------------------

    @staticmethod
    def read(filepath: str, sheet_name: Optional[str] = None) -> List[List[Any]]:
        """Read all data from an Excel sheet into a 2-D list.

        Args:
            filepath:   Path to the .xlsx file.
            sheet_name: Target sheet name. Defaults to the active sheet.

        Returns:
            2-D list of cell values.
        """
        wb = load_workbook(filepath, data_only=True)
        ws = wb[sheet_name] if sheet_name else wb.active
        return [[cell.value for cell in row] for row in ws.iter_rows()]

    @staticmethod
    def get_sheet_names(filepath: str) -> List[str]:
        """Return a list of sheet names in an Excel workbook.

        Args:
            filepath: Path to the .xlsx file.

        Returns:
            List of sheet name strings.
        """
        wb = load_workbook(filepath, read_only=True)
        return wb.sheetnames

    # ------------------------------------------------------------------
    # Writing / Generation
    # ------------------------------------------------------------------

    @staticmethod
    def create(
        filepath: str,
        data: List[List[Any]],
        sheet_name: str = "Sheet1",
        header: Optional[List[str]] = None,
        auto_width: bool = True,
    ) -> str:
        """Create a new Excel workbook and write data to it.

        Args:
            filepath:   Destination path for the .xlsx file.
            data:       2-D list of rows to write.
            sheet_name: Name of the worksheet.
            header:     Optional header row written in bold before *data*.
            auto_width: Automatically adjust column widths to fit content.

        Returns:
            Absolute path of the created file.
        """
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        header_fill = PatternFill("solid", fgColor="4F81BD")
        header_font = Font(bold=True, color="FFFFFF")

        start_row = 1
        if header:
            for col_idx, col_name in enumerate(header, start=1):
                cell = ws.cell(row=1, column=col_idx, value=col_name)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")
            start_row = 2

        for row_idx, row_data in enumerate(data, start=start_row):
            for col_idx, value in enumerate(row_data, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)

        if auto_width:
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    try:
                        max_len = max(max_len, len(str(cell.value or "")))
                    except Exception:
                        pass
                ws.column_dimensions[col_letter].width = min(max_len + 4, 60)

        Path(filepath).parent.mkdir(parents=True, exist_ok=True)
        wb.save(filepath)
        return os.path.abspath(filepath)

    @staticmethod
    def append_rows(
        filepath: str,
        rows: List[List[Any]],
        sheet_name: Optional[str] = None,
        output_path: Optional[str] = None,
    ) -> str:
        """Append rows to an existing Excel sheet.

        Args:
            filepath:    Path to the existing .xlsx file.
            rows:        List of rows (each row is a list of values).
            sheet_name:  Target sheet. Defaults to the active sheet.
            output_path: Save destination. Defaults to overwriting *filepath*.

        Returns:
            Absolute path of the saved file.
        """
        wb = load_workbook(filepath)
        ws = wb[sheet_name] if sheet_name else wb.active
        for row in rows:
            ws.append(row)
        dest = output_path or filepath
        Path(dest).parent.mkdir(parents=True, exist_ok=True)
        wb.save(dest)
        return os.path.abspath(dest)

    # ------------------------------------------------------------------
    # Conversion
    # ------------------------------------------------------------------

    @staticmethod
    def to_csv(
        filepath: str,
        output_path: Optional[str] = None,
        sheet_name: Optional[str] = None,
        encoding: str = "utf-8-sig",
    ) -> str:
        """Convert an Excel sheet to a CSV file.

        Args:
            filepath:    Path to the source .xlsx file.
            output_path: Destination .csv path. Defaults to same stem as source.
            sheet_name:  Source sheet name. Defaults to the active sheet.
            encoding:    File encoding (default utf-8-sig for Excel compatibility).

        Returns:
            Absolute path of the created .csv file.
        """
        data = Excel.read(filepath, sheet_name=sheet_name)
        dest = output_path or Path(filepath).with_suffix(".csv")
        Path(dest).parent.mkdir(parents=True, exist_ok=True)
        with open(dest, "w", newline="", encoding=encoding) as f:
            writer = csv.writer(f)
            writer.writerows(data)
        return os.path.abspath(str(dest))

    @staticmethod
    def from_csv(
        csv_path: str,
        output_path: Optional[str] = None,
        sheet_name: str = "Sheet1",
        has_header: bool = True,
        encoding: str = "utf-8-sig",
    ) -> str:
        """Convert a CSV file to an Excel workbook.

        Args:
            csv_path:    Path to the source .csv file.
            output_path: Destination .xlsx path. Defaults to same stem as source.
            sheet_name:  Name of the target worksheet.
            has_header:  Treat the first row as a header (styled accordingly).
            encoding:    Encoding of the CSV file.

        Returns:
            Absolute path of the created .xlsx file.
        """
        with open(csv_path, newline="", encoding=encoding) as f:
            reader = csv.reader(f)
            rows = list(reader)

        header = rows[0] if has_header and rows else None
        data = rows[1:] if has_header and rows else rows
        dest = output_path or Path(csv_path).with_suffix(".xlsx")
        return Excel.create(str(dest), data, sheet_name=sheet_name, header=header)
