"""Word document processing utilities."""

from __future__ import annotations

import os
from pathlib import Path
from typing import List, Optional

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


class Word:
    """Utilities for creating and manipulating Word (.docx) documents."""

    # ------------------------------------------------------------------
    # Reading
    # ------------------------------------------------------------------

    @staticmethod
    def extract_text(filepath: str) -> str:
        """Extract all plain text from a .docx file.

        Args:
            filepath: Path to the Word document.

        Returns:
            A single string containing all paragraph text joined by newlines.
        """
        doc = Document(filepath)
        return "\n".join(para.text for para in doc.paragraphs)

    @staticmethod
    def get_paragraphs(filepath: str) -> List[str]:
        """Return a list of non-empty paragraph strings from a .docx file.

        Args:
            filepath: Path to the Word document.

        Returns:
            List of paragraph text strings (empty paragraphs are excluded).
        """
        doc = Document(filepath)
        return [para.text for para in doc.paragraphs if para.text.strip()]

    # ------------------------------------------------------------------
    # Writing / Generation
    # ------------------------------------------------------------------

    @staticmethod
    def create(
        filepath: str,
        title: Optional[str] = None,
        paragraphs: Optional[List[str]] = None,
        font_size: int = 12,
    ) -> str:
        """Create a new Word document.

        Args:
            filepath: Destination path for the new .docx file.
            title:    Optional document title (rendered as a heading).
            paragraphs: Optional list of paragraph strings to add.
            font_size:  Default font size in points.

        Returns:
            Absolute path of the created file.
        """
        doc = Document()

        if title:
            doc.add_heading(title, level=1)

        if paragraphs:
            for text in paragraphs:
                para = doc.add_paragraph(text)
                for run in para.runs:
                    run.font.size = Pt(font_size)

        Path(filepath).parent.mkdir(parents=True, exist_ok=True)
        doc.save(filepath)
        return os.path.abspath(filepath)

    @staticmethod
    def add_table(
        filepath: str,
        data: List[List[str]],
        output_path: Optional[str] = None,
        has_header: bool = True,
    ) -> str:
        """Append a table to an existing Word document.

        Args:
            filepath:    Path to the existing .docx file.
            data:        2-D list; first row is treated as header when
                         *has_header* is True.
            output_path: Save destination. Defaults to overwriting *filepath*.
            has_header:  Whether the first row should be bold (header row).

        Returns:
            Absolute path of the saved file.
        """
        doc = Document(filepath)
        if not data:
            doc.save(output_path or filepath)
            return os.path.abspath(output_path or filepath)

        rows, cols = len(data), len(data[0])
        table = doc.add_table(rows=rows, cols=cols, style="Table Grid")

        for r_idx, row_data in enumerate(data):
            for c_idx, cell_text in enumerate(row_data):
                cell = table.cell(r_idx, c_idx)
                cell.text = str(cell_text)
                if has_header and r_idx == 0:
                    for run in cell.paragraphs[0].runs:
                        run.bold = True

        dest = output_path or filepath
        Path(dest).parent.mkdir(parents=True, exist_ok=True)
        doc.save(dest)
        return os.path.abspath(dest)

    # ------------------------------------------------------------------
    # Conversion
    # ------------------------------------------------------------------

    @staticmethod
    def to_txt(filepath: str, output_path: Optional[str] = None) -> str:
        """Convert a Word document to a plain-text file.

        Args:
            filepath:    Path to the source .docx file.
            output_path: Destination .txt path. Defaults to same stem as source.

        Returns:
            Absolute path of the created .txt file.
        """
        text = Word.extract_text(filepath)
        dest = output_path or Path(filepath).with_suffix(".txt")
        Path(dest).parent.mkdir(parents=True, exist_ok=True)
        Path(dest).write_text(text, encoding="utf-8")
        return os.path.abspath(str(dest))
